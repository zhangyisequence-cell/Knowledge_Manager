import io
from zipfile import ZipFile

import openpyxl
import pytest
from backend.adapters.spreadsheet import SpreadsheetAdapter
from backend.config import AppConfig
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula


def save(workbook, path):
    workbook.save(path)
    workbook.close()
    return path


def test_understated_dimensions_do_not_discard_cells(tmp_path):
    book = Workbook()
    book.active['A1'] = 'visible'
    book.active['C3'] = 'must-preserve'
    path = save(book, tmp_path / 'bounds.xlsx')
    output = io.BytesIO()
    with ZipFile(path) as source, ZipFile(output, 'w') as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'ref="A1:C3"', b'ref="A1:A1"')
            target.writestr(entry, data)
    path.write_bytes(output.getvalue())
    text = SpreadsheetAdapter(AppConfig())._read(path)
    assert 'must-preserve' in text
    assert '| 3 |  |  | must-preserve |' in text


@pytest.mark.parametrize('formula,expected', [
    (ArrayFormula(ref='B1:B2', text='=A1:A2*2'), ['=A1:A2*2', 'B1:B2']),
    (DataTableFormula(ref='B1:C3', r1='A1', r2='A2', dt2D=True),
     ['数据表', 'B1:C3', 'A1', 'A2']),
])
def test_special_formulas_preserve_expression_or_table_metadata(tmp_path, formula, expected):
    book = Workbook()
    book.active['A1'] = 2
    book.active['A2'] = 3
    book.active['B1'] = formula
    path = save(book, tmp_path / 'formula.xlsx')
    text = SpreadsheetAdapter(AppConfig())._read(path)
    for part in expected:
        assert part in text
    assert 'object at 0x' not in text


def test_rejected_workbook_closes_archives_before_error_is_released(tmp_path, monkeypatch):
    book = Workbook()
    book.active['A1'] = 1
    book.create_sheet('second')['A1'] = 2
    path = save(book, tmp_path / 'budget.xlsx')
    opened = []
    load = openpyxl.load_workbook

    def tracking_load(*args, **kwargs):
        result = load(*args, **kwargs)
        opened.append(result)
        return result

    monkeypatch.setattr(openpyxl, 'load_workbook', tracking_load)
    adapter = SpreadsheetAdapter(AppConfig())
    adapter.MAX_CELLS = 1
    with pytest.raises(ValueError, match='超过解析上限') as failure:
        adapter._read(path)
    assert failure.value is not None
    assert opened and all(book._archive.fp is None for book in opened)


def test_invalid_ooxml_initialization_closes_underlying_file(tmp_path, monkeypatch):
    import openpyxl.reader.excel

    book = Workbook()
    book.active['A1'] = 'content'
    path = save(book, tmp_path / 'broken.xlsx')
    output = io.BytesIO()
    with ZipFile(path) as source, ZipFile(output, 'w') as target:
        for entry in source.infolist():
            if entry.filename != 'xl/workbook.xml':
                target.writestr(entry, source.read(entry.filename))
    path.write_bytes(output.getvalue())
    opened = []
    archive_type = openpyxl.reader.excel.ZipFile

    def tracking_archive(*args, **kwargs):
        result = archive_type(*args, **kwargs)
        opened.append(result)
        return result

    monkeypatch.setattr(openpyxl.reader.excel, 'ZipFile', tracking_archive)
    with pytest.raises(ValueError, match='Excel') as failure:
        SpreadsheetAdapter(AppConfig())._read(path)
    assert failure.value is not None
    assert opened and all(archive.fp is None or archive.fp.closed for archive in opened)
