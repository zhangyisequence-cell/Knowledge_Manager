"""Real API/file pipeline checks. No model responses are mocked into success."""
from __future__ import annotations

import io
import re
import shutil
import threading
import time
from datetime import date
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZipFile

import pytest
import yaml
from backend.adapters.pdf import PDFAdapter
from backend.config import AIConfig, AppConfig, get_config
from backend.main import app
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from reportlab.pdfgen import canvas


@pytest.fixture
def client(tmp_path, monkeypatch):
    config = AppConfig(
        data_dir=tmp_path / "data",
        database_path=tmp_path / "data" / "jobs.sqlite",
        vault_dir=tmp_path / "vault",
        ai=AIConfig(enabled=False),
        qmd_command="nonexistent-qmd-for-validation",
        max_download_mb=1,
    )
    config.prepare()
    monkeypatch.setattr("backend.main.get_config", lambda: config)
    with TestClient(app) as api:
        yield api, config


def completed(api, response):
    assert response.status_code == 202, response.text
    job_id = response.json()["id"]
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        job = api.get(f"/api/jobs/{job_id}").json()
        if job["status"] in {"succeeded", "failed"}:
            return job
        time.sleep(0.02)
    pytest.fail(f"Job did not complete: {job_id}")


def note_from(api, response):
    job = completed(api, response)
    assert job["status"] == "succeeded", job
    return Path(job["note_path"])


def test_chinese_inspiration_reaches_vault_and_database(client):
    api, config = client
    text = "灵感：建立每周复盘机制。每周五整理知识库，保留原始资料与来源。"
    note = note_from(api, api.post("/api/ingest", json={"text": text, "title": "每周复盘"}))
    assert note.is_relative_to(config.vault_dir)
    content = note.read_text("utf-8")
    assert text in content
    metadata = yaml.safe_load(content.split("---", 2)[1])
    assert metadata["category"] == "待分类"  # fallback must not be called AI classification
    assert metadata["source"] == "text"
    assert api.get("/api/items").json()[0]["title"] == "每周复盘"


@pytest.mark.parametrize("suffix", ["txt", "md"])
def test_text_file_upload_preserves_chinese(client, suffix):
    api, _ = client
    text = "验证样本：知识库必须保留中文正文。"
    note = note_from(api, api.post("/api/upload", files={"file": (f"中文.{suffix}", text.encode())}))
    assert text in note.read_text("utf-8")


def docx_bytes(with_table=False):
    document = Document()
    document.add_paragraph("试点文档：先验证输入到笔记的完整链路。")
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "指标"
        table.cell(0, 1).text = "金额"
        table.cell(1, 0).text = "季度预算"
        table.cell(1, 1).text = "12800"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_word_paragraph_upload(client):
    api, _ = client
    note = note_from(api, api.post("/api/upload", files={"file": ("sample.docx", docx_bytes())}))
    assert "先验证输入到笔记的完整链路" in note.read_text("utf-8")


def test_word_table_data_is_not_silently_lost(client):
    api, _ = client
    note = note_from(api, api.post("/api/upload", files={"file": ("table.docx", docx_bytes(True))}))
    content = note.read_text("utf-8")
    assert "季度预算" in content and "12800" in content


def excel_bytes(blank=False, oversized=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "季度预算"
    if oversized:
        sheet.cell(1048576, 16384, "远端单元格")
    elif not blank:
        sheet.append(["项目", "金额", "日期", "确认"])
        sheet.append(["资料采购", 12800.5, date(2026, 9, 18), False])
        sheet.append(["免费材料", 0, None, True])
        sheet.append(["合计", "=SUM(B2:B3)"])
        workbook.create_sheet("说明").append(["来源|核对\n不能丢失"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.parametrize("suffix", ["xlsx", "xlsm"])
def test_excel_sheets_values_formula_and_original_reach_vault(client, suffix):
    api, config = client
    original = excel_bytes()
    note = note_from(api, api.post("/api/upload", files={"file": (f"预算.{suffix}", original)}))
    text = note.read_text("utf-8")
    for expected in ("季度预算", "说明", "12800.5", "2026-09-18", "FALSE", "TRUE",
                     "SUM(B2:B3)", "未缓存", "来源\\|核对<br>不能丢失"):
        assert expected in text
    assert "| 3 | 免费材料 | 0 |" in text
    metadata = yaml.safe_load(text.split("---", 2)[1])
    assert metadata["source"] == "spreadsheet"
    assert next(config.vault_dir.rglob(f"*.{suffix}")).read_bytes() == original


@pytest.mark.parametrize("case", ["blank", "corrupt", "oversized"])
def test_excel_invalid_input_fails_without_note(client, case):
    api, config = client
    payload = b"not an Excel file" if case == "corrupt" else excel_bytes(
        blank=case == "blank", oversized=case == "oversized"
    )
    job = completed(api, api.post("/api/upload", files={"file": ("invalid.xlsx", payload)}))
    assert job["status"] == "failed"
    assert "Excel" in job["error"]
    assert not list(config.vault_dir.rglob("*.md"))


def test_excel_legacy_xls_preserves_sheets_dates_and_attachment(client):
    api, config = client
    original = (Path(__file__).parent / "fixtures" / "legacy.xls").read_bytes()
    note = note_from(api, api.post("/api/upload", files={"file": ("旧版.xls", original)}))
    text = note.read_text("utf-8")
    for expected in ("预算", "说明", "资料采购", "12800.5", "2026-09-18", "FALSE"):
        assert expected in text
    assert next(config.vault_dir.rglob("*.xls")).read_bytes() == original


def test_excel_cached_formula_result_is_labelled_and_formula_preserved(client):
    api, _ = client
    output = io.BytesIO()
    with ZipFile(io.BytesIO(excel_bytes())) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                assert b"<f>SUM(B2:B3)</f><v></v>" in data
                data = data.replace(b"<f>SUM(B2:B3)</f><v></v>",
                                    b"<f>SUM(B2:B3)</f><v>12800.5</v>")
            target.writestr(entry, data)
    note = note_from(api, api.post("/api/upload", files={"file": ("cached.xlsx", output.getvalue())}))
    text = note.read_text("utf-8")
    assert "公式：=SUM(B2:B3)" in text
    assert "缓存结果（未重新计算）：12800.5" in text


def pdf_bytes(blank=False):
    output = io.BytesIO()
    pdf = canvas.Canvas(output)
    if not blank:
        for row in range(5):
            pdf.drawString(40, 780 - 20 * row, f"Knowledge validation record {row}: preserve original documents and source references.")
    pdf.showPage()
    pdf.save()
    return output.getvalue()


def test_text_pdf_upload(client):
    api, _ = client
    note = note_from(api, api.post("/api/upload", files={"file": ("sample.pdf", pdf_bytes())}))
    assert "Knowledge validation record 4" in note.read_text("utf-8")


def test_empty_pdf_does_not_report_success(client):
    api, _ = client
    job = completed(api, api.post("/api/upload", files={"file": ("empty.pdf", pdf_bytes(True))}))
    assert job["status"] == "failed", "A page heading alone is not extracted document content"


@pytest.mark.parametrize("ocr_text", ["", "## 第 1 页\n\n"])
def test_empty_pdf_guard_independent_of_ocr_installation(client, monkeypatch, ocr_text):
    api, _ = client
    # Simulate OCR that found no text, not a failed OCR executable.
    monkeypatch.setattr(PDFAdapter, "_ocr", lambda self, path: (ocr_text, []))
    job = completed(api, api.post("/api/upload", files={"file": ("blank.pdf", pdf_bytes(True))}))
    assert job["status"] == "failed"
    assert "PDF 未提取到正文" in job["error"]


def test_attachments_are_portable_with_the_vault(client):
    api, config = client
    original = docx_bytes()
    note = note_from(api, api.post("/api/upload", files={"file": ("sample.docx", original)}))
    content = note.read_text("utf-8")
    assert "file:///" not in content, "Host-only file URLs break when the vault is synced"
    assert list(config.vault_dir.rglob("*.docx")), "Original must travel with the vault"
    links = [unquote(link) for link in re.findall(r"\]\(([^)]+)\)", content) if link.endswith(".docx")]
    assert links
    moved_vault = config.vault_dir.parent / "moved-vault"
    shutil.copytree(config.vault_dir, moved_vault)
    moved_note = moved_vault / note.relative_to(config.vault_dir)
    for link in links:
        asset = (moved_note.parent / link).resolve()
        assert asset.is_relative_to(moved_vault)
        assert asset.read_bytes() == original


def test_long_attachment_name_keeps_its_extension(client):
    api, config = client
    original = docx_bytes()
    note_from(api, api.post("/api/upload", files={"file": ("a" * 90 + ".docx", original)}))
    assets = list(config.vault_dir.rglob("*.docx"))
    assert len(assets) == 1
    assert assets[0].read_bytes() == original


def test_unsupported_file_type_fails_without_creating_note(client):
    api, config = client
    job = completed(api, api.post("/api/upload", files={"file": ("sample.unknown", b"unsupported") }))
    assert job["status"] == "failed"
    assert "不支持" in job["error"]
    assert not list(config.vault_dir.rglob("*.md"))


def test_oversize_upload_is_rejected_and_partial_file_removed(client):
    api, config = client
    response = api.post("/api/upload", files={"file": ("large.txt", b"a" * (1024 * 1024 + 1))})
    assert response.status_code == 413
    assert not list((config.data_dir / "originals" / "uploads").glob("*"))


def test_ambiguous_input_is_rejected(client):
    api, _ = client
    assert api.post("/api/ingest", json={"url": "https://example.com", "text": "x"}).status_code == 422


def test_webpage_over_real_local_http(client):
    api, _ = client
    marker = "PUBLICARTICLEVALIDATION42"
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            body = ("<html><head><title>Validation article</title></head><body><article>"
                    + f"<p>{marker}: Preserve sources, review notes weekly, and keep local backups.</p>" * 8
                    + "</article></body></html>").encode()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write(body)
        def log_message(self, *_):
            pass
    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        url = f"http://127.0.0.1:{server.server_port}/article"
        note = note_from(api, api.post("/api/ingest", json={"url": url}))
        content = note.read_text("utf-8")
        assert marker in content and url in content
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_configuration_can_be_isolated_from_real_vault(tmp_path, monkeypatch):
    monkeypatch.setenv("KNOWLEDGE_CONFIG", str(tmp_path / "config.yaml"))
    monkeypatch.setenv("KNOWLEDGE_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("OBSIDIAN_VAULT_DIR", str(tmp_path / "vault"))
    monkeypatch.setenv("AI_ENABLED", "false")
    get_config.cache_clear()
    try:
        config = get_config()
        assert config.vault_dir == tmp_path / "vault"
        assert not config.ai.enabled
    finally:
        get_config.cache_clear()
