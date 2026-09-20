"""Submit synthetic material to a running pilot; save local evidence, never secrets."""
from __future__ import annotations

import argparse
import io
import json
import time
from pathlib import Path

import httpx
from docx import Document
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--url", default="http://127.0.0.1:8787")
    parser.add_argument("--web-url", help="Optional public webpage; requires network")
    args = parser.parse_args()
    results = []
    with httpx.Client(base_url=args.url, timeout=30, trust_env=False) as client:
        health = client.get("/api/health")
        health.raise_for_status()
        def wait(response, label):
            response.raise_for_status()
            job_id = response.json()["id"]
            deadline = time.monotonic() + 120
            while time.monotonic() < deadline:
                status = client.get(f"/api/jobs/{job_id}")
                status.raise_for_status()
                job = status.json()
                if job["status"] in {"succeeded", "failed"}:
                    results.append({"sample": label, **job})
                    print(f"{label}: {job['status']}")
                    return
                time.sleep(0.25)
            raise TimeoutError(f"Timed out: {label}")

        wait(client.post("/api/ingest", json={
            "title": "每周知识复盘（合成验证样本）",
            "text": "每周五整理知识库。为每条笔记保存原始材料与来源，避免只留下无法核对的摘要。",
        }), "Chinese inspiration")
        document = Document()
        document.add_paragraph("知识库试点：验证 Word 正文、表格及附件随库保存。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "指标", "金额"
        table.cell(1, 0).text, table.cell(1, 1).text = "季度预算", "12800"
        output = io.BytesIO()
        document.save(output)
        wait(client.post("/api/upload", files={"file": ("validation.docx", output.getvalue())}), "Word with table")
        output = io.BytesIO()
        pdf = canvas.Canvas(output)
        for row in range(5):
            pdf.drawString(40, 780 - row * 20, "Local knowledge: preserve original sources and review summaries weekly.")
        pdf.save()
        wait(client.post("/api/upload", files={"file": ("validation.pdf", output.getvalue())}), "Text PDF")
        if args.web_url:
            wait(client.post("/api/ingest", json={"url": args.web_url}), "Public webpage")
    report = ROOT / "runtime" / "smoke-results.json"
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps({"health": health.json(), "results": results}, ensure_ascii=False, indent=2), "utf-8")
    return 0 if all(row["status"] == "succeeded" for row in results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
