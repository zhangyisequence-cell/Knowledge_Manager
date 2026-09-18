"""Submit seven synthetic samples; verify terminal jobs and local server artifacts.

Run on the ingestion server: its API exposes note paths, not note contents.
Only HTTP ingestion writes to the application. This script never opens its DB
or changes settings. Evidence and generated inputs stay under --output-dir.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote

import httpx
import yaml
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def samples() -> list[dict]:
    chinese = "每周五整理知识库。负责人李明保留原始材料与来源，周一核对资料采购预算。"
    doc = Document()
    doc.add_paragraph("合成验收文档：负责人李明负责资料归档。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "项目", "金额"
    table.cell(1, 0).text, table.cell(1, 1).text = "资料采购", "12800"
    doc_bytes = io.BytesIO()
    doc.save(doc_bytes)
    workbook = Workbook()
    budget = workbook.active
    budget.title = "预算"
    budget.append(["项目", "金额"])
    budget.append(["资料采购", 12800])
    budget.append(["复核合计", "=SUM(B2:B2)"])
    workbook.create_sheet("说明").append(["合成验收数据", "负责人李明"])
    excel_bytes = io.BytesIO()
    workbook.save(excel_bytes)
    workbook.close()
    pdf_bytes = io.BytesIO()
    pdf = canvas.Canvas(pdf_bytes)
    for i, line in enumerate([
        "Synthetic server acceptance: preserve original sources.",
        "Quarterly budget: 12800. Review archive every Friday.",
        "This PDF contains selectable text, not a scanned image.",
    ]):
        pdf.drawString(40, 780 - i * 20, line)
    pdf.save()
    return [
        {"sample": "chinese_text", "text": chinese,
         "expected": ["负责人李明", "周一核对资料采购预算"]},
        {"sample": "markdown", "filename": "acceptance.md",
         "data": "# 合成验收\n\n周五复盘：保留原始材料。\n- 待办：核对归档清单。\n".encode(),
         "expected": ["周五复盘", "核对归档清单"]},
        {"sample": "plain_text", "filename": "acceptance.txt",
         "data": "合成验收纯文本：负责人李明。采购预算为12800元，周一完成复核。".encode(),
         "expected": ["负责人李明", "12800", "周一完成复核"]},
        {"sample": "word_table", "filename": "acceptance.docx", "data": doc_bytes.getvalue(),
         "expected": ["负责人李明", "资料采购", "12800", "| 项目 | 金额 |"]},
        {"sample": "excel", "filename": "acceptance.xlsx", "data": excel_bytes.getvalue(),
         "expected": ["工作表：预算", "工作表：说明", "资料采购", "12800",
                      "=SUM(B2:B2)", "未缓存，未计算", "负责人李明"]},
        {"sample": "text_pdf", "filename": "acceptance.pdf", "data": pdf_bytes.getvalue(),
         "expected": ["preserve original sources", "12800", "every Friday"]},
        {"sample": "legacy_word", "filename": "legacy.doc",
         "data": (ROOT / "tests" / "fixtures" / "legacy.doc").read_bytes(),
         "expected": ["负责人李明", "资料采购", "12800", "2026-09-30", "周五完成归档核对"]},
    ]


def verify_artifacts(job: dict, sample: dict, ai_enabled: bool, output: Path) -> dict:
    if not job.get("item_id") or not job.get("note_path"):
        raise ValueError("Succeeded job lacks item_id/note_path")
    note = Path(job["note_path"])
    if not note.is_absolute() or not note.is_file():
        raise ValueError("Server note is not locally readable; run this check on the server")
    note_bytes = note.read_bytes()
    text = note_bytes.decode("utf-8").replace("\r\n", "\n")
    if not text.startswith("---\n") or "\n---\n" not in text[4:]:
        raise ValueError("Note has no YAML frontmatter")
    metadata = yaml.safe_load(text.split("---", 2)[1])
    if not isinstance(metadata, dict) or metadata.get("content_id") != job["item_id"]:
        raise ValueError("Note content_id does not match this job")
    original = text.split("# 原始内容\n", 1)
    if len(original) != 2:
        raise ValueError("Note lacks original-content section")
    original_body = original[1].split("\n## 原始附件", 1)[0]
    checks = {value: value in original_body for value in sample["expected"]}
    mode = "ai" if ai_enabled else "rules"
    checks["analysis_mode_matches_health"] = metadata.get("analysis_mode") == mode
    if not ai_enabled:
        checks["explicit_rules_label"] = "规则提取（未启用 AI）" in text
    attachments = []
    if "data" in sample:
        hrefs = re.findall(r"\[附件 \d+\]\(([^)]+)\)", text)
        for href in hrefs:
            relative = Path(unquote(href))
            attachment = (note.parent / relative).resolve()
            if relative.is_absolute() or not attachment.is_relative_to(note.parent.resolve()):
                raise ValueError("Original attachment link is not contained in the note folder")
            stored = attachment.read_bytes()
            attachments.append({"href": href, "path": str(attachment),
                                "sha256": digest(stored), "bytes": len(stored)})
        checks["original_attachment_matches_input"] = any(
            attachment["sha256"] == digest(sample["data"]) for attachment in attachments
        )
        upload = Path(job.get("payload", {}).get("path", ""))
        checks["server_upload_matches_input"] = (
            upload.is_absolute() and upload.is_file()
            and digest(upload.read_bytes()) == digest(sample["data"])
        )
    (output / f"{sample['sample']}-note.md").write_bytes(note_bytes)
    return {"passed": all(checks.values()), "checks": checks,
            "analysis_mode": metadata.get("analysis_mode"), "note_sha256": digest(note_bytes),
            "attachments": attachments, "note_content": text}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default="http://127.0.0.1:8787")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--timeout", type=float, default=180, help="Seconds per job (1-1800)")
    args = parser.parse_args()
    if not 1 <= args.timeout <= 1800:
        parser.error("--timeout must be between 1 and 1800 seconds")
    output = args.output_dir.resolve()
    output.mkdir(parents=True, exist_ok=True)
    report = {"started_at": datetime.now(timezone.utc).isoformat(), "url": args.url,
              "scope": "Synthetic ingestion/storage acceptance; not AI-quality or WeChat acceptance",
              "content_access": "API exposes paths only; notes and attachments read locally",
              "results": [], "passed": False}
    report_path = output / "server-ingestion-results.json"

    def save():
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    save()
    try:
        cases = samples()
        with httpx.Client(base_url=args.url.rstrip("/"), timeout=30, trust_env=False) as client:
            response = client.get("/api/health")
            response.raise_for_status()
            report["health"] = response.json()
            if (report["health"].get("status") != "ok"
                    or report["health"].get("storage_configured") is not True
                    or type(report["health"].get("ai_enabled")) is not bool):
                raise ValueError("Health response does not confirm configured storage and AI mode")
            for sample in cases:
                result = {"sample": sample["sample"], "passed": False, "status": "not_submitted"}
                report["results"].append(result)
                try:
                    title = "合成服务器验收：" + sample["sample"]
                    if "data" in sample:
                        (output / sample["filename"]).write_bytes(sample["data"])
                        result["input_sha256"] = digest(sample["data"])
                        response = client.post("/api/upload", data={"title": title},
                                               files={"file": (sample["filename"], sample["data"])})
                    else:
                        (output / "chinese_text.txt").write_text(sample["text"], encoding="utf-8")
                        response = client.post("/api/ingest",
                                               json={"title": title, "text": sample["text"]})
                    response.raise_for_status()
                    job = response.json()
                    result["job"] = job
                    result["status"] = job["status"]
                    save()
                    deadline = time.monotonic() + args.timeout
                    while job["status"] not in {"succeeded", "failed"}:
                        remaining = deadline - time.monotonic()
                        if remaining <= 0:
                            raise TimeoutError(f"Job did not finish within {args.timeout} seconds")
                        time.sleep(min(0.5, remaining))
                        response = client.get(f"/api/jobs/{job['id']}", timeout=min(30, remaining))
                        response.raise_for_status()
                        job = response.json()
                        result.update(job=job, status=job["status"])
                    if job["status"] == "succeeded":
                        result.update(verify_artifacts(job, sample, report["health"]["ai_enabled"], output))
                    else:
                        result["error"] = job.get("error") or "Ingestion job failed"
                except (httpx.HTTPError, OSError, ValueError, KeyError, TypeError, yaml.YAMLError) as error:
                    result["error"] = f"{type(error).__name__}: {error}"
                    if isinstance(error, TimeoutError):
                        result["timed_out"] = True
                save()
                print(f"{sample['sample']}: {result['status']}; verified={result['passed']}")
        report["passed"] = len(report["results"]) == 7 and all(
            result["passed"] for result in report["results"]
        )
    except (httpx.HTTPError, OSError, ValueError, KeyError, TypeError, yaml.YAMLError) as error:
        report["error"] = f"{type(error).__name__}: {error}"
    finally:
        report["finished_at"] = datetime.now(timezone.utc).isoformat()
        save()
    print(f"Evidence: {report_path}")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
